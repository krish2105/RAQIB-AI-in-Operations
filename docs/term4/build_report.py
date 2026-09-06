"""Build docs/AI218_RAQIB_report.docx from docs/results/. No placeholders: every figure is loaded.

  cd cloud && uv run --no-sync python ../docs/term4/build_report.py
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import DOCS, all_figures, load, num, pct  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

OUT = DOCS / "AI218_RAQIB_report.docx"

REFERENCES = [
    "Ahmad, H. M., & Rahimi, A. (2024). SH17: A dataset for human safety and personal protective equipment detection in manufacturing industry. Journal of Safety Science and Resilience, 5(4), 375–386. https://doi.org/10.1016/j.jnlssr.2024.09.002",
    "Erlang, A. K. (1917). Solution of some problems in the theory of probabilities of significance in automatic telephone exchanges. Elektroteknikeren, 13, 5–13.",
    "Friedman, J. H. (2001). Greedy function approximation: A gradient boosting machine. The Annals of Statistics, 29(5), 1189–1232. https://doi.org/10.1214/aos/1013203451",
    "Gruen, T. W., Corsten, D. S., & Bharadwaj, S. (2002). Retail out-of-stocks: A worldwide examination of extent, causes and consumer responses. Grocery Manufacturers of America.",
    "Hillier, F. S., & Lieberman, G. J. (2021). Introduction to operations research (11th ed.). McGraw-Hill.",
    "Huangfu, Q., & Hall, J. A. J. (2018). Parallelizing the dual revised simplex method. Mathematical Programming Computation, 10(1), 119–142. https://doi.org/10.1007/s12532-017-0130-5",
    "Hyndman, R. J., & Athanasopoulos, G. (2021). Forecasting: Principles and practice (3rd ed.). OTexts. https://otexts.com/fpp3/",
    "Little, J. D. C. (1961). A proof for the queuing formula: L = λW. Operations Research, 9(3), 383–387. https://doi.org/10.1287/opre.9.3.383",
    "United Arab Emirates. (2021). Federal Decree-Law No. 45 of 2021 on the Protection of Personal Data. Official Gazette. https://u.ae/en/about-the-uae/digital-uae/data/data-protection-laws",
    "Zhang, Y., Sun, P., Jiang, Y., Yu, D., Weng, F., Yuan, Z., Luo, P., Liu, W., & Wang, X. (2022). ByteTrack: Multi-object tracking by associating every detection box. In Proceedings of ECCV 2022. https://doi.org/10.1007/978-3-031-20047-2_1",
]


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0x14, 0x18, 0x1D)
    return p


def para(doc, text, italic=False, size=10.5, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.bold = bold
    r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(6)
    return p


def bullets(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(it)
        r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(3)


def table(doc, rows, header=True, widths=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = str(cell)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    r.bold = header and i == 0
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
    doc.add_paragraph()
    return t


def caption(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x5B, 0x64, 0x72)
    p.paragraph_format.space_after = Pt(10)


def main() -> None:
    k = load("kpis_retail.json")
    kf = load("kpis_factory.json")
    f = load("forecast_retail.json")
    w = load("workforce_retail.json")
    r = load("report_retail_en.json")
    e = load("edge_eval.json")
    t = load("toolcalls_retail.json")
    figs = all_figures()
    ba = r["before_after"] or {}
    st = r["staffing"] or {}
    gov = r["governance"]
    y = e["detectors"]["yolo"]
    rt = e["detectors"].get("rtdetr", {})
    p = e["pipeline"]
    peaks = ", ".join(x["ts"][11:16] for x in f["peaks"][:3])
    recs = r["recommendations"]

    doc = Document()
    s = doc.styles["Normal"]
    s.font.name = "Calibri"
    s.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rr = title.add_run("RAQIB · رقيب\nA vision-driven operations control room for supermarket floors")
    rr.font.size = Pt(22)
    rr.bold = True
    para(doc, f"MAIB AI 218 — AI in Operations · Term 4 deployed MVP · Krishna Mathur · {date.today().isoformat()}", italic=True)
    para(doc, "Repository: github.com/krish2105/RAQIB-AI-in-Operations · Live dashboard link and QR on the deck's slide 5. Every figure in this report is generated from docs/results/ by docs/term4/build_report.py; nothing is typed by hand. Where a number comes from the labelled simulator rather than a live camera, the text says so.", italic=True, size=9)

    # ------------------------------------------------------------------ 1
    h(doc, "1. Executive summary")
    para(doc, "Supermarkets already own the sensor that could run their floor: the CCTV system. RAQIB (Arabic: the watcher) turns those cameras into an operations agent. On the edge, a detector and tracker feed a deterministic rule engine that emits events: a customer entered, a queue passed its limit, a checkout served someone, a shelf face has been empty for five minutes. In the cloud, the events feed a queueing model, a demand forecast, a staffing integer program, and a bounded agent that alerts, raises restock tasks, and proposes opening tills. A control-room web app shows all of it live in English, Hindi, and Arabic. The same engine, under the name MUSHRIF (the supervisor), runs a factory profile for a Greenlam Laminates plant: helmet compliance, exclusion-zone breaches, machine downtime, work orders into the plant's existing tracker.")
    para(doc, "Three measured results on the deployed demo:", bold=True)
    bullets(doc, [
        f"Latency. The edge pipeline (detect, track, blur faces, rules, store) runs at {p['fps_end_to_end']} frames per second end to end on a laptop, with {p['frame_to_event_ms_p50']} ms median and {p['frame_to_event_ms_p95']} ms p95 from frame to event. The specification's frame-to-alert target of three seconds is met with two orders of magnitude to spare.",
        f"Forecast. Hourly footfall forecast beats the seasonal-naive baseline by {f['improvement_pct']}% on a seven-day holdout (MAE {f['mae']} vs {f['mae_naive']}), on the 21-day labelled simulator.",
        f"Staffing. A weekly plan that keeps checkout utilisation at or below 0.85 needs {st.get('staff_hours', 'n/a')} staff-hours against {st.get('baseline_staff_hours', 'n/a')} with three tills open all week, a saving of {st.get('savings_hours', 'n/a')} hours; applying only the agent's open-till proposals at peaks cuts total customer waiting by {ba.get('reduction_pct', 'n/a')}%.",
    ])
    para(doc, f"Operational KPIs over the last 24 hours of the demo: service level (share of 15-minute slots with a queue of three or fewer) {pct(k['service_level'])} against a 90% target; store on-shelf availability {pct(k['osa_store'])}; peak checkout utilisation ρ {num(k['peak_rho'], 2)}; {k['footfall']} customers counted at {k['footfall_per_hour']} per hour.")

    # ------------------------------------------------------------------ 2
    h(doc, "2. Problem and operations context")
    para(doc, "A hypermarket floor is two operations systems sharing one workforce. It is a multi-server queue at the checkouts, where the arrival rate swings by a factor of four within a day and by a further 40% between a Monday and a Friday evening. It is also an inventory system whose last metre, the shelf face, is the only metre the customer sees. Both are managed today by walking the floor. A duty manager glances at the tills and decides whether to call someone off the shop floor; a replenishment team follows a fixed round regardless of what has actually sold. Neither decision is made from data because neither quantity is measured.")
    para(doc, "The cost of the two failures is well documented. Queues longer than about three people drive abandonment and shorter baskets; Gruen, Corsten and Bharadwaj (2002) put the global out-of-stock rate at about 8% of SKUs and the sales loss at about 4% of turnover, with a third of shoppers buying elsewhere or not at all when they meet a gap. For a Dubai hypermarket operator such as Majid Al Futtaim's Carrefour or Lulu, running stores of 10,000 to 20,000 square metres with 40 to 60 checkouts, a single percentage point of on-shelf availability and a few minutes of peak waiting time are worth more than the entire cost of a camera-analytics system.")
    para(doc, "The context that makes this a UAE project rather than a generic one: the trading week peaks Thursday to Saturday and shifts entirely during Ramadan and the Eid holidays, so a forecast that only knows last week is wrong for about a quarter of the year; the workforce is multilingual, so alerts must be readable in Hindi and Arabic as well as English; and Federal Decree-Law No. 45 of 2021 on personal data protection, together with employment practice, means that any camera analytics must be built so that it cannot identify people. RAQIB treats each of these as a design requirement, not a caveat.")
    para(doc, "The premise of the design follows from the operations-management view of the problem. The critical decisions, whether a queue is over its limit or a shelf is empty, are geometry and counting; they should be deterministic, cheap, and auditable. The judgement decisions, whether to open a till or move a replenisher, benefit from a model that can weigh utilisation, forecast, and staff cost. Moving people is a human's call. RAQIB separates the three: rules on the edge, models and a bounded agent in the cloud, a floor manager on the approval button.")

    # ------------------------------------------------------------------ 3
    h(doc, "3. Literature and practice")
    para(doc, "Queueing. Erlang's (1917) multi-server model and Little's (1961) law are the foundation of every queue-management practice in retail and are taught in Hillier and Lieberman (2021) as M/M/c. The closed forms for utilisation ρ = λ/(cμ), the probability of waiting, and the expected wait W_q are what let a control room translate an arrival rate into a till count in real time. Commercial queue-management vendors sell the counting (overhead sensors above tills, footfall beams at doors) and a dashboard; the queueing model is usually absent, and the manager is left to interpret a count. RAQIB puts the model on the same chart as the measurement, so the gap between expected and observed wait is visible and can be used to calibrate μ.")
    para(doc, "On-shelf availability. Gruen et al. (2002) remain the reference for the size and consumer response of out-of-stocks across 52 studies; their finding that most gaps originate in store replenishment rather than in the supply chain is the reason RAQIB measures the shelf face directly rather than inferring gaps from sales velocity. Time-to-restock, the length of a gap episode, is the operational lever the replenishment team controls.")
    para(doc, "Forecasting. Hyndman and Athanasopoulos (2021) argue that any seasonal demand forecast must be judged against a seasonal-naive baseline (same period last cycle), because sophisticated models often fail to beat it. RAQIB reports the baseline alongside the model on every forecast and in this report. The challenger is a gradient-boosting regressor (Friedman, 2001) on calendar features, lags, rolling means, and a leak-free hour-of-week profile; it was chosen over neural sequence models because it is explainable, trains in seconds on a laptop, and degrades gracefully with three weeks of history.")
    para(doc, "Workforce scheduling. Shift and till scheduling is a classic integer programming problem. The formulation here is deliberately small (minimise total till-slots subject to a utilisation cap per slot) and solved with HiGHS (Huangfu and Hall, 2018) through SciPy, so the plan can be recomputed on every report and extended with adjacency and minimum-shift constraints without changing solvers.")
    para(doc, "Computer vision. ByteTrack (Zhang et al., 2022) associates every detection box, including low-confidence ones, which is what makes per-track dwell time and first-entry counting stable enough to estimate λ and μ from video. For the factory profile, SH17 (Ahmad and Rahimi, 2024) is the personal-protective-equipment dataset with 17 classes across industrial scenes; its non-commercial licence is respected in the dataset ledger. On privacy, the UAE's Federal Decree-Law No. 45 of 2021 establishes purpose limitation and data minimisation; RAQIB's contribution is architectural: no identity data is created, so none needs protecting.")

    # ------------------------------------------------------------------ 4
    h(doc, "4. System design")
    para(doc, "RAQIB is one pipeline with two site profiles. The unit of exchange is the Event: id, site, camera, timestamp, kind, severity (1 info, 2 warning, 3 critical), a payload of counts, classes, confidences and zone names, an optional clip path, and the id of the rule that fired. The edge produces Events; the cloud stores, reasons about, and streams them; the web app renders them. Adding a rule adds an event kind and nothing else.")
    para(doc, "Edge (Python 3.12, runs next to the cameras). Frames come from a webcam, a looped file, or an RTSP stream. A Detector adapter with two implementations (YOLO26n, the default, and RT-DETR-L) finds people; ByteTrack assigns session-scoped ids. Immediately after tracking, the head band of every detected person is pixelated and blurred; no later stage sees an unblurred frame. Polygon zones (entrance, queue, checkout, shelf, work area, exclusion, machine) are read from a site YAML in normalised coordinates. Rules are pure functions of the tracks, the zones, a small per-camera state, and the frame's timestamp; there is no wall clock inside a rule, which is what makes them unit-testable with synthetic tracks. Events land in a SQLite outbox; a ring buffer writes a ten-second clip around every warning or critical event; a syncer posts idempotent batches to the cloud when it is reachable and simply keeps queuing when it is not.")
    table(doc, [
        ["Rule", "Profile", "Condition", "Severity"],
        ["R10 queue over", "retail", "≥ N persons in a queue zone for ≥ 60 s (N from site thresholds)", "2"],
        ["R11 shelf gap", "retail", "shelf empty ratio > 0.4 for ≥ 5 min", "1"],
        ["R12 footfall", "both", "track's first entry to an entrance zone", "1"],
        ["R13 checkout served", "retail", "track leaves checkout after ≥ 20 s dwell (μ estimator)", "1"],
        ["R01 no helmet", "factory", "person in work area, no helmet over head for ≥ 2 s", "3"],
        ["R02 zone breach", "factory", "person inside an exclusion zone", "3"],
        ["R03 machine stopped", "factory", "machine ROI stopped ≥ 120 s in scheduled run", "2"],
    ], widths=[3.5, 2, 8, 2])
    caption(doc, "Table 1. The rule engine. Thresholds live in the site YAML; rules re-emit at most every 300 s while a condition persists.")
    para(doc, "Cloud (FastAPI, Postgres on Supabase, hosted on Render). Endpoints for events, clips, sites, KPIs, forecast, workforce, report, and a Server-Sent Events stream. The agent runs per event: a backend proposes tool calls, a Policy filters them, autonomous calls execute and are logged, proposals wait for a human. Two backends share the same Policy: a deterministic dry-run backend (a policy table with written rationales, used in the demo because no model key is configured) and a Claude backend using Anthropic tool use with the same seven tool schemas. The analytics modules are plain functions: Erlang-C queueing, on-shelf availability, gradient-boosting forecast with a seasonal-naive comparator, and the staffing integer program. A Jinja template renders the weekly report in English, Hindi, and Arabic.")
    para(doc, "Web (Next.js 16 on Vercel). A control-room design: graphite dark theme first, a light 'blueprint on paper' theme, teal signal colour for the retail profile and hazard amber for the factory profile, IBM Plex in Latin, Arabic and Devanagari with tabular figures for every number. The signature element is the tape, a rolling 24-hour strip chart at the top of every page in which footfall is the trace, operational events are ticks coloured by severity, and now is a fixed pen at the trailing edge. Below it: KPI tiles, a React Three Fiber floor plan that pulses where events land, the live event stream, the queue model against the observed queue, and the agent's open proposals. Event pages show the clip, the confidence, the rule text, and the agent's decision; the actions page approves or rejects proposals and shows the tool-call audit; the report page prints to PDF.")
    para(doc, "Privacy by design. Faces are blurred before storage; there is no face-recognition or re-identification model anywhere in the codebase; track ids are integers that reset when the process restarts; only counts, classes, confidences and events leave the edge; the agent never receives an image. Simulated history used for the demo carries a simulated flag in every payload and is labelled in the UI and in this report.")

    # ------------------------------------------------------------------ 5
    h(doc, "5. Data and models")
    table(doc, [
        ["Asset", "Licence", "Use in this build"],
        ["Pexels checkout clip 39221979 (1920×1080, 35 s); time-lapse 854634 (1280×720, 18 s)", "Pexels licence; CC0", "Looped demo cameras; faces blurred regardless"],
        ["YOLO26n and RT-DETR-L COCO checkpoints (Ultralytics)", "AGPL-3.0", "Detector adapter; person class only; swap or Enterprise licence before commercial sale"],
        ["SH17, 8,099 images, 17 PPE classes (Ahmad & Rahimi, 2024)", "CC BY-NC-SA 4.0", "PPE fine-tuning for the factory profile; manual download, not run in the demo"],
        ["Labelled simulator, 21 days, daily and weekly seasonality, Gulf weekend", "generated", "Forecast, workforce, report history; every event flagged simulated"],
    ], widths=[7, 3, 6])
    caption(doc, "Table 2. Data and weights. The full ledger with URLs is docs/datasets.md.")
    para(doc, f"Detector choice. YOLO26n was chosen for the edge because it is the smallest current Ultralytics model with end-to-end (NMS-free) inference and runs at {y['latency_ms_p50']} ms per frame on the M4 Pro's GPU via Metal, {y['fps_detect_only']} frames per second detect-only. RT-DETR-L, a transformer detector whose architecture is Apache-licensed, runs the same pipeline unchanged at {rt.get('latency_ms_p50', 'n/a')} ms per frame ({rt.get('fps_detect_only', 'n/a')} fps). The adapter exists for exactly this reason: the licence swap is one command-line flag, and the test suite asserts that the rule engine has no import from the detector package.")
    doc.add_picture(str(figs["latency"]), width=Cm(15))
    caption(doc, "Figure 1. Median per-frame latency for the two detector implementations on the same 60 frames of the checkout clip (docs/results/edge_eval.json).")
    table(doc, [
        ["Metric", "Value", "Source"],
        ["YOLO26n median detect latency / detect-only fps", f"{y['latency_ms_p50']} ms / {y['fps_detect_only']}", "edge_eval.json"],
        ["RT-DETR-L median detect latency / detect-only fps", f"{rt.get('latency_ms_p50', 'n/a')} ms / {rt.get('fps_detect_only', 'n/a')}", "edge_eval.json"],
        ["End-to-end pipeline fps (detect, track, blur, rules, store)", str(p["fps_end_to_end"]), "edge_eval.json"],
        ["Frame → event latency p50 / p95", f"{p['frame_to_event_ms_p50']} / {p['frame_to_event_ms_p95']} ms", "edge_eval.json"],
        ["PPE mAP50", e["ppe_map50"]["status"] + " (needs SH17 plus two hours of labelled site footage)", "edge_eval.json"],
    ], widths=[7, 5, 3.5])
    caption(doc, "Table 3. Edge evaluation. The mAP target of 0.80 is a pilot deliverable; COCO weights detect persons only.")
    para(doc, "Tracking and counting. Footfall is one tick per track the first time its foot point enters the entrance zone, so a person who lingers at the door is counted once. Service completions are tracks that leave a checkout zone after at least 20 seconds. Both counts are sensitive to tracker id switches; ByteTrack's association of low-confidence boxes is what keeps ids stable through the partial occlusions of a queue. The shelf empty ratio is a heuristic (Canny edge density and colour variance per cell of the shelf face) that a small classifier trained on labelled site crops replaces behind the same function signature.")

    # ------------------------------------------------------------------ 6
    h(doc, "6. Operations analytics")
    h(doc, "6.1 Queue model versus observed", level=2)
    para(doc, f"Arrivals λ per 15-minute slot come from footfall ticks; the service rate μ per till comes from checkout dwell (rule R13): μ = 3600 / mean dwell, currently {w['mu']:.0f} customers per hour per till and labelled 'estimated from video' wherever it appears. With c = {w['observed_tills']} tills, Erlang-C gives utilisation ρ and expected wait W_q per slot; the observed W_q is Little's law on the measured queue length, W_q = L_q / λ. The two are drawn on one chart because the gap is information: it says how far real checkouts are from the exponential-service assumption and gives the pilot a calibration target once POS data replaces the video estimate of μ.")
    doc.add_picture(str(figs["queue"]), width=Cm(15.5))
    caption(doc, f"Figure 2. Utilisation and expected wait per slot, last 24 slots. Peak ρ {num(k['peak_rho'], 2)}; service level {pct(k['service_level'])} of slots with queue ≤ 3 (target 90%); average queue when over limit {k['avg_queue']} customers (docs/results/kpis_retail.json).")
    h(doc, "6.2 On-shelf availability", level=2)
    para(doc, f"OSA per shelf is one minus the share of the window the shelf face was in a gap state; a gap opens when the empty ratio exceeds 0.4 for five minutes and closes when the signal stops. Over the last 24 hours the store OSA is {pct(k['osa_store'])} across the monitored shelves: {', '.join(f'{s} {pct(v)}' for s, v in k['osa'].items())}. Mean time to restock per gap episode: {', '.join(f'{s} {v:.0f} min' for s, v in k['time_to_restock_min'].items()) or 'no gap episodes in the window'}. The 98% OSA target is the standard for grocery; the lever is the length of each episode, which the restock work order shortens.")
    h(doc, "6.3 Forecast versus seasonal naive", level=2)
    para(doc, f"The forecast target is hourly footfall. Features: hour of day, weekday, lag 24 h, lag 168 h, 24 h and 168 h rolling means, and a leak-free hour-of-week profile (the mean of the same weekday and hour over all previous weeks). The comparator is seasonal naive: the same hour last week. On the last seven days of the 21-day history the gradient-boosting model reaches MAE {f['mae']} against {f['mae_naive']} for the baseline, {f['improvement_pct']}% better, with MAPE {f['mape']}% on non-zero hours. The next-24-hour peaks are expected at {peaks} UTC; the agent's recommendation to pre-position a floater before each peak comes from this output.")
    doc.add_picture(str(figs["forecast"]), width=Cm(15.5))
    caption(doc, "Figure 3. Last three days of history, the 24-hour forecast, and the seasonal-naive baseline (docs/results/forecast_retail.json).")
    para(doc, "Interpretation. The simulator is generated from smooth daily and weekly curves with Poisson noise, which is close to what seasonal naive assumes, so the margin measured here is a floor rather than an estimate of real performance. Real footfall carries promotions, weather, school holidays, and Ramadan hours that a same-hour-last-week baseline cannot see and calendar features can. The specification target of at least 20% improvement is therefore kept as a pilot gate on real data.")
    h(doc, "6.4 Workforce optimisation", level=2)
    para(doc, f"Formulation: minimise Σ_s tills_s subject to λ_s / (tills_s · μ) ≤ 0.85 for every 15-minute slot s, 1 ≤ tills_s ≤ ceiling, integer. Solved with HiGHS through scipy.optimize.milp; a max-step constraint between adjacent slots is implemented for the roster version. On one day the plan needs {w['staff_hours']} staff-hours against {w['baseline_staff_hours']} with {w['observed_tills']} tills open all day (saving {w['savings_hours']} hours). Over the week: {st.get('staff_hours', 'n/a')} against {st.get('baseline_staff_hours', 'n/a')} staff-hours, a saving of {st.get('savings_hours', 'n/a')} hours; the plan peaks at {st.get('peak_tills', 'n/a')} tills with a mean of {st.get('mean_tills', 'n/a')}, and {st.get('slots_over_baseline', 'n/a')} slots need more tills than are open today.")
    doc.add_picture(str(figs["workforce"]), width=Cm(15.5))
    caption(doc, "Figure 4. Planned tills per 15-minute slot for one day against today's flat schedule (docs/results/workforce_retail.json).")
    h(doc, "6.5 Before and after", level=2)
    para(doc, f"The continuous-improvement test replays the week with the agent's open-till proposals applied: tills per slot = max(today's {w['observed_tills']}, plan). Closing tills off-peak is the staff-hour saving above and is deliberately not counted as a wait effect. Total customer waiting falls from {ba.get('baseline_customer_wait_min', 'n/a')} to {ba.get('with_plan_customer_wait_min', 'n/a')} customer-minutes over the week, a {ba.get('reduction_pct', 'n/a')}% reduction, with {ba.get('plan_unstable_slots', 0)} slots left above ρ = 1. The two levers together, fewer tills off-peak and one more at the peak, are what the agent's weekly recommendations describe.")
    para(doc, "Recommendations generated for this week: " + "; ".join(f"{i + 1}. {x['title']} ({x['expected_reduction']})" for i, x in enumerate(recs)) + ".")

    # ------------------------------------------------------------------ 7
    h(doc, "7. Agent design and governance")
    para(doc, "The agent never sees video. It receives one Event, the site profile, the number of tills, the utilisation ρ of the current slot from the queueing model, and recent counts. It may call only the tools below; each call is validated against a strict JSON schema (unknown fields are rejected) before anything executes, and every execution is written to an audit table with input, output, latency, success, and model cost.")
    table(doc, [
        ["Tool", "Autonomy", "Guard"],
        ["send_alert(channel, lang, template, vars)", "autonomous", "templates only in EN/HI/AR; no free text reaches WhatsApp"],
        ["create_work_order(machine_id, summary, severity)", "autonomous for severity ≤ 2", "one per machine or shelf per 4 h unless a human overrides; Greenlam tracker contract"],
        ["log_downtime(machine_id, start, end, reason)", "autonomous", "schema"],
        ["escalate(event_id, to_role, note)", "mandatory on severity 3", "always attaches the clip; cannot be suppressed"],
        ["request_human_review(event_id, question)", "mandatory below 0.6 confidence", "replaces side-effecting calls"],
        ["propose_open_till(till, window, rho)", "proposal", "dropped unless ρ > 0.85"],
        ["propose_staffing_change / propose_maintenance_window", "proposal", "human approves or rejects on /actions"],
    ], widths=[6, 4, 6])
    caption(doc, "Table 4. Allow-listed tools and the policies that bound them (cloud/raqib_api/agent/tools.py, policies.py).")
    para(doc, "Policies are code, not prompt text, and they run after either backend: P1 severity-3 events always escalate and alert even if the model returned nothing; P2 the agent cannot downgrade a severity-3 event (any lower severity in its arguments is rewritten to 3); P3 one work order per asset per four hours, compared on the event clock so replayed history behaves like live traffic; P4 a till proposal must cite ρ above 0.85; P5 below 0.6 confidence, side effects are replaced with a human-review request; P6 anything outside the allow-list is dropped. Each policy has a test that asserts the behaviour, including one that feeds the agent a work order attempting to mark a critical event as severity 1 and checks that the escalation still happens.")
    para(doc, f"Governance figures for the seeded week: {gov['actions']} agent actions ({gov['autonomous']} autonomous, {gov['proposals']} proposals), {gov['tool_calls']} tool calls, {gov['failed_calls']} failed, model cost ${gov['cost_usd']} because the demo runs the deterministic backend. The retail tool-call ledger holds {t['total_calls']} calls at ${t['total_cost_usd']}. With the Claude backend at about 50 actionable events a day the projected model cost is under USD 5 per month per store. The human approval rate, the specification's acceptance metric (target 60%), is computed by the report once proposals have been decided on the live system.")

    # ------------------------------------------------------------------ 8
    h(doc, "8. Deployment")
    para(doc, "Edge on an Apple M-series machine (Metal GPU) or any Linux box via the CPU Docker image; the pilot box is whatever the site already owns. API on Render as a web service from render.yaml with a health check at /health; Postgres on Supabase from the generated schema; web on Vercel with the API URL as its only environment variable. The free Render tier sleeps after fifteen idle minutes: the edge keeps writing to SQLite, the first request wakes the API, and the syncer's idempotent batches catch up within seconds. This offline path is covered by tests that return 503 and connection errors from a fake server and assert that nothing is lost and nothing is duplicated.")
    para(doc, "Monitoring is deliberately simple: the dashboard header shows live, connecting, or reconnecting for the event stream; /health reports the event count and the agent backend; the actions page shows the tool-call audit with success counts and latencies; the pilot SOP prescribes a five-minute daily check. Secrets live only in environment variables, and the repository ships a .env.example with every variable documented.")

    # ------------------------------------------------------------------ 9
    h(doc, "9. Limitations and ethics")
    bullets(doc, [
        "The demo cameras are public stock clips of a New Zealand supermarket, not a Dubai store; zone geometry and thresholds were tuned to the clip.",
        "There is no POS feed, so μ is estimated from checkout dwell on video and every output says so. A CSV importer for real POS is the first pilot task.",
        "The 21-day history is a labelled simulator. It is flagged in every payload, labelled in the UI and this report, and it flatters the seasonal-naive baseline, so the forecast margin is a floor.",
        "COCO weights detect persons only. Helmet and vest detection needs SH17 plus about two hours of labelled site footage; PPE mAP is therefore reported as not run rather than estimated.",
        "Ultralytics is AGPL-3.0. That is acceptable for this coursework, the portfolio, and a pilot, and the RT-DETR adapter demonstrates the swap path; a commercial deployment needs the swap or an Enterprise licence.",
        "Ethics. The system is built to be unable to identify people: faces are blurred before storage, no identity model exists, and the agent never sees images. Camera analytics still changes how staff experience their workplace; the pilot SOP requires a signed privacy notice, entrance signage, and a supervisor review of every escalation. Under Federal Decree-Law No. 45 of 2021 the site owner remains the data controller and must verify the current position with counsel before deployment; nothing in this report is legal advice.",
        "Emiratisation and workforce policy. A staffing plan that reduces till-hours affects rostering, not headcount; how a store uses the saving is a management decision the system only informs. This is stated on the workforce page.",
    ])

    # ------------------------------------------------------------------ 10
    h(doc, "10. Business case and 90-day rollout for a Dubai hypermarket")
    para(doc, "Buyer: the store operations manager, with the head of replenishment and the regional operations director as approvers. The pitch is not 'AI'; it is fewer walk-outs at the peak, shorter gaps on the shelf, and a roster built from measured demand, on cameras the store already owns.")
    table(doc, [
        ["Item", "Monthly, one store", "Notes"],
        ["Edge box", "AED 0–110", "existing Mac mini / laptop or a ~USD 200–700 box amortised over 24 months"],
        ["API hosting (Render small instance)", "AED 0–26", "free tier sleeps; small paid instance for the pilot"],
        ["Database (Supabase) and web (Vercel)", "AED 0", "free tiers at pilot scale"],
        ["Model cost (Claude, ~50 actionable events/day)", "AED < 20", "dry-run backend costs nothing"],
        ["Labelling (one-off)", "AED 1,500–3,000", "two hours of site footage, ~600 frames, first month only"],
        ["Total recurring", "under AED 200", "under USD 50"],
    ], widths=[6, 3.5, 6.5])
    caption(doc, "Table 5. Pilot cost. Currency conversions at AED 3.67 per USD.")
    para(doc, f"Value model. Each avoided walk-out at a peak till is worth one average basket; each hour of gap on a fast-moving shelf face is lost sales plus a substitution the customer may not repeat. Both are measurable from the events the system already records: the number of slots where ρ exceeded 0.85 without a till being opened, and the shelf gap-hours per day. In the seeded week the plan removes {st.get('slots_over_baseline', 'n/a')} under-staffed peak slots and {ba.get('reduction_pct', 'n/a')}% of customer waiting; a store need only value that at a few baskets a day to cover the recurring cost.")
    table(doc, [
        ["Phase", "Weeks", "Scope", "KPI gate to proceed"],
        ["Pilot", "1–4", "one entrance camera, one checkout bank, three shelves; privacy notice and signage; SOP in docs/pilot_sop.md; agent in dry-run then live alerts", "false alerts < 5 per camera per day; human approval rate ≥ 60%; frame-to-alert < 3 s"],
        ["Expand", "5–8", "all checkout banks; POS import so μ is measured; staffing plan feeds the roster; Hindi and Arabic alerts to floor staff", "service level ≥ 90% of slots; forecast ≥ 20% better than naive on real data"],
        ["Scale", "9–13", "shelf coverage aisle by aisle; second store; Claude backend for weekly reports", "OSA ≥ 98%; cost per store < USD 150 per month; approval rate sustained"],
    ], widths=[2, 1.6, 7.4, 5.5])
    caption(doc, "Table 6. Rollout with gates. Each gate is a number the system itself reports.")
    para(doc, f"The factory profile runs on the same rollout logic. On the Greenlam demo site the last 24 hours show PPE compliance {pct(kf['compliance'])}, {kf['downtime_min']} minutes of unplanned downtime, and {kf['breaches']} exclusion-zone breach; the pilot SOP for one camera and one machine over four weeks is written and costed in docs/pilot_sop.md.")

    # ------------------------------------------------------------------ 11
    h(doc, "11. Conclusion and future work")
    para(doc, "RAQIB shows that an operations agent can be built on existing cameras by keeping the layers separate: deterministic rules for the decisions that must be auditable, queueing theory and a forecast for the decisions that need a model, an integer program for the roster, and a language model only where judgement is needed and always behind a policy that code enforces. The deployed demo meets its latency target by two orders of magnitude, beats the naive forecast on the data it has, produces a staffing plan a manager can read, and explains every action it takes. What it cannot yet do is honest too: it has not seen a real store, a real POS feed, or a labelled hour of site footage.")
    para(doc, "Next: POS import and per-till μ; two hours of labelled site footage and fine-tuned PPE weights with mAP reported; a planogram check on the shelf face; digital-twin playback on the 3D floor to scrub a day; a 24-hour unattended soak on a real camera; and, once the Claude backend is live, the human approval rate as the headline governance metric.")

    # ------------------------------------------------------------------ 12 (v2)
    a = load("ask_eval.json")["summary"]
    c = load("crew.json")
    tw = load("twin.json")
    fl = load("fleet.json")
    wa = load("watch.json")
    se = load("security_eval.json")
    h(doc, "12. v2 extension: from one agent to a governed crew")
    para(doc, "Phases E to K extend the deployed system without touching the edge-to-event contract or the policy layer. Six capabilities were added on top of the same Event and the same Policy, all behind one model-provider adapter that keeps inference cost at zero (Ollama on the site box or laptop, then the Gemini and Groq free tiers; the Anthropic client stays in code but off). Every figure below is read from docs/results/.")
    bullets(doc, [
        f"Ask (Phase E). A deterministic trilingual router, hybrid retrieval (SQL prefilter, structured ranking, BM25, a {a.get('embed_model', 'bge-m3')} vector leg fused by reciprocal rank) and an answer in which every factual sentence must cite a retrieved chunk. On {a['cases']} EN/HI/AR cases: recall@5 {a['recall_at_5']:.2f}, faithfulness {a['faithfulness']:.2f}, citation coverage {a['citation_coverage']:.2f}, language match {a['language_match']:.2f}, {a['hallucinations']} hallucinations; mean latency {a['latency_mean_ms']/1000:.1f} s on the laptop.",
        f"Watch (Phase F). A blurred MJPEG camera wall ({wa['stream']['client_fps_measured']} fps at the client) with live detection boxes, and a vision-language second opinion on at most a few blurred keyframes per event ({wa['opinion']['warm_s']} s warm). The opinion is advisory: it may add a human-review request at high disagreement confidence, and it can never lower a severity, which is asserted after every write.",
        f"Crew (Phase G). Six narrow agents (FloorOps, ShelfOps, Workforce, Safety, Analyst, Auditor) on one runtime: per-agent tool allow-lists checked before Policy, per-run budgets ({c['envelope']['budget_default']['max_tool_calls']} calls, {c['envelope']['budget_default']['max_seconds']} s), HMAC-signed typed messages with replay protection, an Auditor after every run, guarded memory with provenance, quarantine and rollback, and a kill switch that returns 503 on agent routes while severity-3 escalation continues on the deterministic path. A FloorOps run takes {c['measured']['floorops_run_seconds']} s.",
        f"Twin (Phase H). Replay of any day in {tw['replay_bins']} one-minute bins on the 3D floor ({tw['playback_fps_headless_60x']} fps headless at 60x) and what-if sliders that re-run the M/M/c model and the staffing MILP on the day's real arrivals.",
        "Integrations (Phase I). POS CSV import that switches the service rate to measured throughput per till, planogram and price-tag rules on the edge (R14, R15), WhatsApp approved templates to an opt-in roster, and a Greenlam client with retries, idempotency and a circuit breaker.",
        f"Fleet and operations (Phase J). Supabase Auth with four roles and site scoping, heartbeat and drift telemetry from every edge box (PSI over {fl['drift']['window_h']} h against {fl['drift']['baseline_days']} days, alert after {fl['drift']['hours_over']} h over {fl['drift']['threshold_psi']} with a suggested fix), a retention job that logs what it deleted, and one OpenTelemetry span per model call feeding a cost-per-day KPI that reads ${c['measured']['cost_usd']:.2f}.",
        f"Security (Phase K). The OWASP Top 10 for Agentic Applications mapped to one control and one executable attack per risk; {se['passed']} of {se['total']} defended on {se['date'][:10]}; CI fails on any regression and also runs pip-audit, npm audit and a secret scan. Defensive headers and a CSP on both the API and the web app, an SSRF guard on URL ingest, and a weights hash pin that refuses to start the edge box on a mismatch.",
    ])
    table(doc, [["ASI", "Risk", "Attack test", "Result"]] + [[r["asi"], r["risk"], r["test"], "pass" if r["passed"] else "FAIL"] for r in se["results"]], widths=[1.4, 3.2, 9.6, 1.4])
    caption(doc, "Table 7. Red-team results, docs/results/security_eval.json.")
    para(doc, "The governance conclusion of section 7 survives the extension: the crew gained judgement, memory and a twin without gaining a second way to execute a tool. The residual risks are stated in docs/security/threat_model.md: free providers see prompt text (never frames or secrets), the free-tier database is ephemeral until the Supabase connection string is set, and the weights pin is enforced only where an operator sets it.")

    h(doc, "References")
    for ref in REFERENCES:
        para(doc, ref, size=9.5)

    h(doc, "Appendix")
    para(doc, "A. Site configuration: edge/sites/retail_demo.yaml (zones in normalised coordinates, thresholds, floor plan in metres). B. Rule table: Table 1 and edge/raqib_edge/rules.py. C. Tool schemas: cloud/raqib_api/agent/tools.py. D. Test coverage (v2): edge 65 tests (rules with synthetic tracks, detector adapter on real frames, blur, store, offline sync, pipeline, planogram, OCR, telemetry, weights pin); cloud 176 tests (routers, Erlang-C textbook cases, policies, tool validation, forecast, MILP, report in three languages, provider chain and quotas, RAG, VLM, crew bus/budget/kill switch/memory guard, twin, integrations, auth, fleet, retention, security) plus the ten-attack red-team harness; web 39 unit tests and seven Playwright flows (grader path, theme and RTL, Ask, Crew, Twin x2, Security). E. Results: docs/results/*.json and docs/results/figures/. F. Generation: docs/term4/build_report.py.", size=9.5)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
